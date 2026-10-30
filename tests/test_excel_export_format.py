import os
import openpyxl
from playwright.sync_api import sync_playwright

def test_excel_export_format():
    print("Testing Excel export format via Playwright...")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(accept_downloads=True)
        page = context.new_page()

        console_errors = []
        page.on("console", lambda msg: console_errors.append(msg.text) if msg.type == "error" else None)

        page.goto("http://127.0.0.1:5000", wait_until="networkidle")

        # Click the Excel Export button dropdown
        btn_excel_main = page.locator("#btnExcelMain")
        assert btn_excel_main.is_visible(), "Excel button must be visible"
        btn_excel_main.click()

        # Listen for download
        with page.expect_download() as download_info:
            page.locator("#exportAllExcel").click()

        download = download_info.value
        download_path = os.path.join("exports", "test_downloaded.xlsx")
        download.save_as(download_path)
        print(f"Downloaded file saved to: {download_path}")

        # Open and inspect the downloaded Excel workbook
        wb = openpyxl.load_workbook(download_path)
        print("Workbook sheets:", wb.sheetnames)
        assert wb.sheetnames == ["Jobs", "Resume Tips"], f"Expected ['Jobs', 'Resume Tips'], got {wb.sheetnames}"

        # 1. Verify 'Jobs' sheet
        ws_jobs = wb["Jobs"]
        expected_jobs_headers = ["Job Title", "Company", "Location", "Posted Date", "Apply Link"]
        actual_jobs_headers = [ws_jobs.cell(1, c).value for c in range(1, 6)]
        print("Jobs sheet headers:", actual_jobs_headers)
        assert actual_jobs_headers == expected_jobs_headers, f"Jobs headers mismatch: {actual_jobs_headers}"

        # Verify row 2 data
        title = ws_jobs.cell(2, 1).value
        company = ws_jobs.cell(2, 2).value
        location = ws_jobs.cell(2, 3).value
        posted_date = str(ws_jobs.cell(2, 4).value)
        apply_cell = ws_jobs.cell(2, 5)

        print(f"Row 2: Title={title}, Company={company}, Location={location}, Date={posted_date}, Apply={apply_cell.value}")
        assert title and len(title) > 0, "Title must not be empty"
        assert company and len(company) > 0, "Company must not be empty"
        assert location and len(location) > 0, "Location must not be empty"
        assert len(posted_date) >= 10 and posted_date[4] == '-' and posted_date[7] == '-', f"Invalid date format: {posted_date}"
        assert apply_cell.value == "Apply", f"Expected cell text 'Apply', got {apply_cell.value}"
        assert apply_cell.hyperlink and apply_cell.hyperlink.target.startswith("http"), f"Missing or invalid hyperlink target: {apply_cell.hyperlink}"

        # 2. Verify 'Resume Tips' sheet
        ws_tips = wb["Resume Tips"]
        expected_tips_headers = ["Action", "Recommendation"]
        actual_tips_headers = [ws_tips.cell(1, 1).value, ws_tips.cell(1, 2).value]
        print("Resume Tips headers:", actual_tips_headers)
        assert actual_tips_headers == expected_tips_headers, f"Resume tips headers mismatch: {actual_tips_headers}"

        assert ws_tips.max_row == 16, f"Expected 16 rows (1 header + 15 recommendations), got {ws_tips.max_row}"
        assert ws_tips.cell(2, 1).value == "ADD"
        assert "Azure/AWS AI" in str(ws_tips.cell(2, 2).value)
        assert ws_tips.cell(11, 1).value == "REMOVE"
        assert "Class 12" in str(ws_tips.cell(11, 2).value)

        # Check console errors
        assert len(console_errors) == 0, f"JS console errors occurred: {console_errors}"

        browser.close()
        print("\nALL EXCEL EXPORT FORMAT CHECKS PASSED PERFECTLY!")

if __name__ == "__main__":
    test_excel_export_format()

